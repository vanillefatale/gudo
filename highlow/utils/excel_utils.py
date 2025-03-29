from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl import Workbook

def save_with_format(df, filename):
    """단일 시트를 가진 엑셀 파일 저장 (퍼센트 서식 적용 포함)"""
    wb = Workbook()
    ws = wb.active

    # 데이터 추가
    for r in dataframe_to_rows(df, index=False, header=True):
        ws.append(r)

    # 서식 처리
    _apply_formats(ws, df)

    wb.save(filename)


def save_multiple_sheets(dfs_with_names, filename):
    """여러 시트를 가진 엑셀 파일 저장 (퍼센트 서식 포함)"""
    wb = Workbook()
    wb.remove(wb.active)  # 기본 시트 제거

    for sheet_name, df in dfs_with_names:
        ws = wb.create_sheet(title=sheet_name)

        for r in dataframe_to_rows(df, index=False, header=True):
            ws.append(r)

        _apply_formats(ws, df)

    wb.save(filename)


def _apply_formats(ws, df):
    """퍼센트 서식 적용 (현재는 전일거래량대비율만 처리)"""
    headers = df.columns.tolist()

    if "전일거래량대비율" in headers:
        percent_col = headers.index("전일거래량대비율") + 1
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=percent_col)
            if isinstance(cell.value, (int, float)):
                cell.value = cell.value / 100
                cell.number_format = "0.00%"
